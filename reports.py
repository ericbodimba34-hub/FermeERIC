import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os
import sqlite3
from datetime import datetime

def generate_excel_report(db_path, output_path):
    wb = openpyxl.Workbook()
    
    # Common styles
    header_fill = PatternFill(start_color="1B5E20", end_color="1B5E20", fill_type="solid")
    sub_header_fill = PatternFill(start_color="2E7D32", end_color="2E7D32", fill_type="solid")
    accent_fill = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")
    alert_fill = PatternFill(start_color="FFEBEE", end_color="FFEBEE", fill_type="solid")
    
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Segoe UI", size=14, bold=True, color="1B5E20")
    bold_font = Font(name="Segoe UI", size=10, bold=True)
    normal_font = Font(name="Segoe UI", size=10)
    
    thin_border = Border(
        left=Side(style='thin', color='DDDDDD'),
        right=Side(style='thin', color='DDDDDD'),
        top=Side(style='thin', color='DDDDDD'),
        bottom=Side(style='thin', color='DDDDDD')
    )
    
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    cur = conn.cursor()

    # --- FEUILLE 1 : SYNTHÈSE GLOBALE & FINANCIÈRE ---
    ws1 = wb.active
    ws1.title = "Synthèse Financière"
    ws1.views.sheetView[0].showGridLines = True
    
    ws1['A1'] = "LA FERME DE DIÉGONEFLA — APPLICATION DFM"
    ws1['A1'].font = title_font
    ws1['A2'] = f"Rapport de gestion intégré généré le {datetime.now().strftime('%d/%m/%Y à %H:%M')}"
    ws1['A2'].font = Font(name="Segoe UI", size=10, italic=True, color="555555")
    
    # Financial indicators
    cur.execute("SELECT COALESCE(SUM(montant_net), 0), COALESCE(SUM(montant_regle), 0), COALESCE(SUM(montant_restant), 0) FROM ventes")
    ca_total, ca_encaisse, creances_total = cur.fetchone()
    
    cur.execute("SELECT COALESCE(SUM(montant), 0) FROM depenses")
    depenses_total = cur.fetchone()[0]
    benefice_net = ca_total - depenses_total

    ws1['A4'] = "INDICATEURS FINANCIERS CLÉS"
    ws1['A4'].font = bold_font
    
    kpis = [
        ("Chiffre d'Affaires Total (Net)", f"{ca_total:,.0f} FCFA".replace(',', ' ')),
        ("Total Encaissé", f"{ca_encaisse:,.0f} FCFA".replace(',', ' ')),
        ("Total Créances Clients", f"{creances_total:,.0f} FCFA".replace(',', ' ')),
        ("Total Dépenses & Charges", f"{depenses_total:,.0f} FCFA".replace(',', ' ')),
        ("Bénéfice Net d'Exploitation", f"{benefice_net:,.0f} FCFA".replace(',', ' '))
    ]
    
    r = 5
    for label, val in kpis:
        ws1[f'A{r}'] = label
        ws1[f'B{r}'] = val
        ws1[f'A{r}'].font = bold_font if r == 9 else normal_font
        ws1[f'B{r}'].font = bold_font
        ws1[f'A{r}'].fill = accent_fill if r == 9 else PatternFill(fill_type=None)
        ws1[f'B{r}'].fill = accent_fill if r == 9 else PatternFill(fill_type=None)
        r += 1

    # Détail des Ventes
    r += 2
    ws1[f'A{r}'] = "HISTORIQUE RÉCENT DES VENTES"
    ws1[f'A{r}'].font = bold_font
    r += 1
    
    cols_v = ["N° Facture", "Date", "Client", "Net à Payer (FCFA)", "Réglé (FCFA)", "Solde Dû (FCFA)", "Mode", "Statut"]
    for c_idx, col in enumerate(cols_v, 1):
        cell = ws1.cell(row=r, column=c_idx, value=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    
    cur.execute("SELECT numero_facture, date_vente, client_nom, montant_net, montant_regle, montant_restant, mode_paiement, statut_paiement FROM ventes ORDER BY date_vente DESC")
    for row in cur.fetchall():
        r += 1
        ws1.cell(row=r, column=1, value=row['numero_facture'])
        ws1.cell(row=r, column=2, value=row['date_vente'])
        ws1.cell(row=r, column=3, value=row['client_nom'])
        ws1.cell(row=r, column=4, value=row['montant_net']).number_format = '#,##0 "FCFA"'
        ws1.cell(row=r, column=5, value=row['montant_regle']).number_format = '#,##0 "FCFA"'
        ws1.cell(row=r, column=6, value=row['montant_restant']).number_format = '#,##0 "FCFA"'
        ws1.cell(row=r, column=7, value=row['mode_paiement'].upper())
        ws1.cell(row=r, column=8, value=row['statut_paiement'].upper())
        for c in range(1, 9):
            ws1.cell(row=r, column=c).border = thin_border

    # Détail des Dépenses
    r += 3
    ws1[f'A{r}'] = "HISTORIQUE DES DÉPENSES & CHARGES"
    ws1[f'A{r}'].font = bold_font
    r += 1
    
    cols_d = ["Référence", "Date", "Catégorie", "Libellé", "Fournisseur", "Montant (FCFA)", "Mode Paiement"]
    for c_idx, col in enumerate(cols_d, 1):
        cell = ws1.cell(row=r, column=c_idx, value=col)
        cell.font = header_font
        cell.fill = sub_header_fill
        cell.alignment = Alignment(horizontal="center")
        
    cur.execute("SELECT reference, date_depense, categorie, libelle, fournisseur, montant, mode_paiement FROM depenses ORDER BY date_depense DESC")
    for row in cur.fetchall():
        r += 1
        ws1.cell(row=r, column=1, value=row['reference'])
        ws1.cell(row=r, column=2, value=row['date_depense'])
        ws1.cell(row=r, column=3, value=row['categorie'].upper())
        ws1.cell(row=r, column=4, value=row['libelle'])
        ws1.cell(row=r, column=5, value=row['fournisseur'])
        ws1.cell(row=r, column=6, value=row['montant']).number_format = '#,##0 "FCFA"'
        ws1.cell(row=r, column=7, value=row['mode_paiement'].upper())
        for c in range(1, 8):
            ws1.cell(row=r, column=c).border = thin_border

    # --- FEUILLE 2 : ÉLEVAGE VOLAILLES ---
    ws2 = wb.create_sheet(title="Élevage Volailles")
    ws2.views.sheetView[0].showGridLines = True
    ws2['A1'] = "SUIVI DU CHEPTEL AVICOLE (POULETS)"
    ws2['A1'].font = title_font
    
    cols_vol = ["Code Lot", "Désignation", "Souche", "Date Entrée", "Eff. Initial", "Eff. Actuel", "Poids Moyen (g)", "Bâtiment", "Statut"]
    for c_idx, col in enumerate(cols_vol, 1):
        cell = ws2.cell(row=3, column=c_idx, value=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        
    cur.execute("SELECT code, nom, souche, date_entree, effectif_initial, effectif_actuel, poids_moyen_g, batiment, statut FROM lots_volailles")
    r = 3
    for row in cur.fetchall():
        r += 1
        ws2.cell(row=r, column=1, value=row['code'])
        ws2.cell(row=r, column=2, value=row['nom'])
        ws2.cell(row=r, column=3, value=row['souche'])
        ws2.cell(row=r, column=4, value=row['date_entree'])
        ws2.cell(row=r, column=5, value=row['effectif_initial'])
        ws2.cell(row=r, column=6, value=row['effectif_actuel'])
        ws2.cell(row=r, column=7, value=row['poids_moyen_g'])
        ws2.cell(row=r, column=8, value=row['batiment'])
        ws2.cell(row=r, column=9, value=row['statut'].upper())
        for c in range(1, 10):
            ws2.cell(row=r, column=c).border = thin_border

    # --- FEUILLE 3 : ÉLEVAGE CUNICOLE ---
    ws3 = wb.create_sheet(title="Élevage Lapins")
    ws3.views.sheetView[0].showGridLines = True
    ws3['A1'] = "SUIVI DU CHEPTEL CUNICOLE (LAPINS & REPRODUCTEURS)"
    ws3['A1'].font = title_font
    
    ws3['A3'] = "REPRODUCTEURS INDIVIDUELS"
    ws3['A3'].font = bold_font
    cols_rep = ["Tatouage", "Nom", "Race", "Sexe", "Cage", "Statut", "Total Portées", "Notes"]
    for c_idx, col in enumerate(cols_rep, 1):
        cell = ws3.cell(row=4, column=c_idx, value=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        
    cur.execute("SELECT code_tatouage, nom, race, sexe, cage_numero, statut, portees_total, notes FROM lapins_reproducteurs")
    r = 4
    for row in cur.fetchall():
        r += 1
        ws3.cell(row=r, column=1, value=row['code_tatouage'])
        ws3.cell(row=r, column=2, value=row['nom'])
        ws3.cell(row=r, column=3, value=row['race'])
        ws3.cell(row=r, column=4, value=row['sexe'].upper())
        ws3.cell(row=r, column=5, value=row['cage_numero'])
        ws3.cell(row=r, column=6, value=row['statut'].upper())
        ws3.cell(row=r, column=7, value=row['portees_total'])
        ws3.cell(row=r, column=8, value=row['notes'])
        for c in range(1, 9):
            ws3.cell(row=r, column=c).border = thin_border

    # --- FEUILLE 4 : STOCKS & ALIMENTS ---
    ws4 = wb.create_sheet(title="Stocks & Aliments")
    ws4.views.sheetView[0].showGridLines = True
    ws4['A1'] = "ÉTAT CONSOLIDÉ DES STOCKS & VALORISATION"
    ws4['A1'].font = title_font
    
    cols_stk = ["Code", "Désignation", "Catégorie", "Stock Actuel", "Unité", "Seuil Alerte", "Prix Unitaire (FCFA)", "Valeur Stock (FCFA)", "Statut"]
    for c_idx, col in enumerate(cols_stk, 1):
        cell = ws4.cell(row=3, column=c_idx, value=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        
    cur.execute("SELECT code, designation, categorie, quantite_stock, unite, seuil_alerte, prix_unitaire, (quantite_stock * prix_unitaire) as valeur, statut FROM stocks_articles")
    r = 3
    for row in cur.fetchall():
        r += 1
        ws4.cell(row=r, column=1, value=row['code'])
        ws4.cell(row=r, column=2, value=row['designation'])
        ws4.cell(row=r, column=3, value=row['categorie'].upper())
        ws4.cell(row=r, column=4, value=row['quantite_stock'])
        ws4.cell(row=r, column=5, value=row['unite'])
        ws4.cell(row=r, column=6, value=row['seuil_alerte'])
        ws4.cell(row=r, column=7, value=row['prix_unitaire']).number_format = '#,##0 "FCFA"'
        ws4.cell(row=r, column=8, value=row['valeur']).number_format = '#,##0 "FCFA"'
        st_cell = ws4.cell(row=r, column=9, value=row['statut'].upper())
        if row['quantite_stock'] <= row['seuil_alerte']:
            st_cell.fill = alert_fill
        for c in range(1, 10):
            ws4.cell(row=r, column=c).border = thin_border

    # Auto-adjust column widths for all sheets
    for sheet in wb.worksheets:
        for col in sheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                try:
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                except:
                    pass
            sheet.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

    wb.save(output_path)
    conn.close()
    return output_path

if __name__ == '__main__':
    db = os.path.join(os.path.dirname(__file__), 'data', 'dfm_ferme.db')
    out = os.path.join(os.path.dirname(__file__), 'data', 'rapport_dfm_test.xlsx')
    generate_excel_report(db, out)
    print("Report generated:", out)
